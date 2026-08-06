# -*- coding: utf-8 -*-
"""FASE ZERO — raccolta degli URL dal corpus ITALIA NERA.

Percorre fonti/ e corpus/, estrae gli URL con il loro contesto, classifica il dominio in
sette categorie e scrive repertorio/URL_CENSITI.xlsx (un foglio per categoria, ordinato per
frequenza del dominio). Il controllo HTTP dello stato è in verifica_url.py (passata separata).

Regole (CLAUDE.md §4): i .docx del corpus ereditato sono testo semplice con estensione
impropria → open(utf-8); i .docx veri (firma PK) → python-docx + XML grezzo (per gli
hyperlink nelle relazioni). I PDF si saltano in questa passata."""

import os
import re
import zipfile
from collections import Counter, defaultdict

import openpyxl

URL_RE = re.compile(r'https?://[^\s,;)\]<>"\']+')
# code di punteggiatura da ripulire in coda all'URL
TRAILING = '.,;:)]}>"\'»'

# domini tecnici da escludere (namespace XML, non fonti)
ESCLUSI = ('schemas.openxmlformats.org', 'schemas.microsoft.com', 'purl.org',
           'www.w3.org', 'schemas.openxmlformats', 'openoffice.org', 'w3.org/1999',
           'ns.adobe.com', 'schemas.oasis-open.org')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---- classificazione dei domini in 7 categorie ----
CAT = [
    ("giudiziario o parlamentare", (
        'senato.it','camera.it','parlamento.it','parlamento17.camera','inchieste.camera',
        'gazzettaufficiale.it','giustizia.it','cortedicassazione','curia.europa.eu','hudoc',
        'echr.coe.int','anticorruzione.it','giurcost.org','normattiva','documenti.camera')),
    ("archivio o istituzione pubblica", (
        'archiv','.gov','gov.','beniculturali','cultura.gov.it','memoria.cultura','siusa',
        'bundesarchiv','bstu','nara','stasi','acs.','san.beniculturali','cnr.it','nato.int',
        'europa.eu','coe.int','wilsoncenter','sicurezzanazionale.gov','icar.cultura',
        'diplo.de','bundestag.de','historikerkommission','uni-marburg','clio-online')),
    ("editoria accademica", (
        '.edu','jstor','springer','tandfonline','cambridge.org','oup.com','doi.org',
        'academia.edu','researchgate','sciencedirect','unibo.it','unige.it','unimib.it',
        'sturzo.it','storiadigitale.it','archiviopenale','penaledp.it','ilmondodegliarchivi',
        'siscalt.it','treccani.it','unipd.it')),
    ("enciclopedia collaborativa", (
        'wikipedia.org','wikimedia','wikisource','wikiwand','fandom.com','grokipedia',
        'de-academic','en-academic','vereins.fandom')),
    ("stampa", (
        'corriere.it','repubblica.it','ilfattoquotidiano','ansa.it','today.it','fanpage.it',
        'ilpost.it','editorialedomani','quotidianodelsud','ilrestodelcarlino','unionesarda',
        'tagesspiegel','taz.de','wsws.org','spiegel','theguardian','nytimes','dire.it',
        'reggioreport','reggionline','citynow.it','corrieredellacalabria','gazzettadelsud',
        'zoom24','ilsicilia','focus.it','askanews','fanpage','t-online.de','heise.de',
        'telepolis','opendemocracy','inquirer','yahoo','newsinfo')),
    ("sito militante o di parte", (
        'peacelink','liberainformazione','antimafiaduemila','avvisopubblico','nsu-watch',
        'staatsunrecht','verfassungsblog','terzultimafermata','iacchite','lavialibera',
        'aktion-freiheitstattangst','panopticon.blog','lecorte.de','vorwaerts','referio',
        'amadeu-antonio','dka-kanzlei','federprivacy','giornaledicalabria')),
]

def categoria(dominio):
    d = dominio.lower()
    for nome, chiavi in CAT:
        if any(k in d for k in chiavi):
            return nome
    # regole per suffisso/TLD e forme ricorrenti (riducono «altro»)
    if any(t in d for t in ('.gov', '.go.jp', '.go.kr', '.gc.ca', 'gov.', '.mil',
                            'un.org', 'nato.int', 'europa.eu', 'oecd.org', 'nato')):
        return "archivio o istituzione pubblica"
    if any(t in d for t in ('.edu', '.ac.', 'jstor', 'journal', 'apjjf', 'jstage',
                            'muse.jhu', 'cambridge', 'oxford', 'sciencepo', 'brill')):
        return "editoria accademica"
    if any(t in d for t in ('wiki', 'fandom', 'britannica', 'academic', 'infogalactic',
                            'citizendium')):
        return "enciclopedia collaborativa"
    if any(t in d for t in ('news', 'times', 'post', 'guardian', 'bbc', 'reuters',
                            'asahi', 'japantimes', 'scmp', 'nikkei', 'gazzett', 'giornale',
                            'quotidiano', 'zeitung', 'spiegel', 'welt', 'faz.net')):
        return "stampa"
    if any(t in d for t in ('blog', 'tumblr', 'wordpress', 'blogspot', 'medium.com',
                            'substack', 'libcom', 'marxists', 'anarch', 'sokaglobal',
                            'sgi', 'komei.or', '.or.jp', 'party', 'workers')):
        return "sito militante o di parte"
    return "altro"

def dominio_di(url):
    m = re.match(r'https?://([^/]+)/?', url)
    return (m.group(1).lower() if m else url).lstrip('www.')

def pulisci(url):
    return url.rstrip(TRAILING)

def testo_docx_reale(path):
    """Testo dei paragrafi + XML grezzo di document.xml e delle rels (per gli hyperlink)."""
    parti = []
    try:
        from docx import Document
        d = Document(path)
        parti.append("\n".join(p.text for p in d.paragraphs))
    except Exception:
        pass
    try:
        z = zipfile.ZipFile(path)
        for nome in ('word/document.xml', 'word/_rels/document.xml.rels',
                     'word/footnotes.xml', 'word/_rels/footnotes.xml.rels'):
            if nome in z.namelist():
                parti.append(z.read(nome).decode('utf-8', 'ignore'))
    except Exception:
        pass
    return "\n".join(parti)

def leggi(path):
    low = path.lower()
    if low.endswith('.pdf'):
        return None
    if low.endswith('.docx'):
        with open(path, 'rb') as f:
            sig = f.read(2)
        if sig == b'PK':                      # docx vero
            return testo_docx_reale(path)
        with open(path, encoding='utf-8', errors='ignore') as f:  # docx-testo
            return f.read()
    if low.endswith(('.txt', '.md', '.csv', '.xml', '.html')):
        with open(path, encoding='utf-8', errors='ignore') as f:
            return f.read()
    return None

def main():
    record = []           # (url, dominio, file, categoria, contesto)
    per_file = Counter()
    for base in ('fonti', 'corpus'):
        d = os.path.join(ROOT, base)
        if not os.path.isdir(d):
            continue
        for r, _, files in os.walk(d):
            for fn in sorted(files):
                if fn == '.gitkeep':
                    continue
                path = os.path.join(r, fn)
                testo = leggi(path)
                if not testo:
                    continue
                rel = os.path.relpath(path, ROOT)
                for m in URL_RE.finditer(testo):
                    url = pulisci(m.group(0))
                    dom = dominio_di(url)
                    if any(x in dom for x in ESCLUSI):
                        continue
                    ctx = testo[max(0, m.start()-200):m.start()].replace('\n', ' ').strip()
                    record.append((url, dom, rel, categoria(dom), ctx))
                    per_file[rel] += 1

    # unicità per (url) mantenendo la prima occorrenza + conteggio occorrenze
    visti = {}
    for url, dom, rel, cat, ctx in record:
        if url not in visti:
            visti[url] = [url, dom, rel, cat, ctx, 1]
        else:
            visti[url][5] += 1
    unici = list(visti.values())

    dom_freq = Counter(u[1] for u in unici)
    cat_gruppi = defaultdict(list)
    for u in unici:
        cat_gruppi[u[3]].append(u)

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ordine = [c[0] for c in CAT] + ["altro"]
    for cat in ordine:
        ws = wb.create_sheet(cat[:31])
        ws.append(["URL", "dominio", "categoria", "file di provenienza",
                   "occorrenze", "freq. dominio", "contesto (200 char)",
                   "stato HTTP", "data controllo"])
        righe = sorted(cat_gruppi.get(cat, []),
                       key=lambda u: (-dom_freq[u[1]], u[1], u[0]))
        for url, dom, rel, c, ctx, n in righe:
            ws.append([url, dom, c, rel, n, dom_freq[dom], ctx[:200], "", ""])

    out = os.path.join(ROOT, 'repertorio', 'URL_CENSITI.xlsx')
    wb.save(out)

    tot_occ = len(record)
    tot_uni = len(unici)
    tot_dom = len(dom_freq)
    enc = sum(1 for u in unici if u[3] == "enciclopedia collaborativa")
    parte = sum(1 for u in unici if u[3] == "sito militante o di parte")
    print(f"file con URL: {len(per_file)}")
    print(f"URL totali (occorrenze): {tot_occ}")
    print(f"URL distinti: {tot_uni}  su  {tot_dom} domini distinti")
    print(f"  enciclopedie collaborative (<=C): {enc}")
    print(f"  siti di parte (<=C):              {parte}")
    print(f"  totale <=C: {enc+parte}  ({round(100*(enc+parte)/max(tot_uni,1))}% dei distinti)")
    print("categorie (URL distinti):")
    for cat in ordine:
        print(f"  {len(cat_gruppi.get(cat, [])):>4}  {cat}")
    print("primi 12 domini per frequenza:")
    for dom, n in dom_freq.most_common(12):
        print(f"  {n:>4}  {dom}")
    print(f"\nscritto: {out}")

if __name__ == "__main__":
    main()
